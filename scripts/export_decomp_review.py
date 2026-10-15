"""分解結果（decomp2026.json）を人が確認・承認するための xlsx を出す。
  python3 scripts/export_decomp_review.py <decomp.json> <out.xlsx>
列: No / タイトル（仮） / 競技 / シートの説明 / 読んだ元（映像・画像・説明のみ） / 映像で読んだ事実 / 実況 / 値（強い順、根拠つき） / 確認（OK・要修正） / 修正メモ
"""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
src, out = sys.argv[1], sys.argv[2]
d = json.load(open(src)); moments = {m["no"]: m for m in json.load(open("data/moments2026.json"))}
wb = Workbook(); ws = wb.active; ws.title = "分解の確認"
head = ["No", "タイトル（仮）", "競技", "シートの説明", "読んだ元", "映像で読んだ事実", "実況", "値（強い順・根拠）", "確認", "修正メモ"]
ws.append(head)
for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="111111")
SRC = {"video": "映像＋音声", "image": "画像", "text": "説明のみ"}
for no in sorted(d):
    v = d[no]; m = moments.get(no, {})
    if "error" in v: ws.append([no, m.get("title"), m.get("sportLabel"), m.get("description"), "エラー", v["error"]]); continue
    f = v.get("facts") or {}
    facts = " / ".join(x for x in [f.get("situation"), f.get("score"), "・".join(f.get("people") or []), "・".join(f.get("teams") or [])] if x)
    vals = "\n".join(f"{x['v']}（{x['s']:.1f}）{x['evidence']}" for x in sorted(v["values"], key=lambda x: -x["s"]))
    ws.append([no, m.get("title"), m.get("sportLabel"), m.get("description"), SRC.get(v.get("source"), v.get("source")), facts, "\n".join(v.get("commentary") or []), vals, "", ""])
widths = [6, 22, 14, 40, 10, 36, 30, 44, 8, 24]
for i, w in enumerate(widths): ws.column_dimensions[chr(65 + i)].width = w
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
ws.freeze_panes = "C2"
lex = json.load(open("data/lexicon2026.json"))
ws2 = wb.create_sheet("辞書（画面に出る言葉）")
ws2.append(["軸", "値", "入力の同義語"])
for a in lex["axes"]:
    for v in a["values"]: ws2.append([a["label"], v["v"], "、".join(v["syn"])])
ws2.column_dimensions["A"].width = 8; ws2.column_dimensions["B"].width = 16; ws2.column_dimensions["C"].width = 80
wb.save(out); print("saved", out, len(d))
